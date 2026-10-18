"""Styled PDF + Word renderer for the project documents (MeetingAI look: navy, violet, blue).

Markdown dialect (on top of normal headings / lists / tables / images / code):
    ## 4. Section title            -> numbered section banner (new section)
    LEAD: one sentence             -> shown inside the banner (must follow the ## line)
    ### 4.1 Sub heading            -> sub heading with accent bar
    > [!ASSUMPTION] Optional title -> callout box; kinds: NOTE, ASSUMPTION, DECISION, KEY, CAUTION, TIP
    > text lines ...
    ```stats                       -> row of big-number tiles, lines "value | label"
    ![Caption](images/x.png)       -> numbered figure
    <pagebreak>                    -> force a new page
"""
from __future__ import annotations

import io
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from PIL import Image as PILImage  # noqa: E402

from core import mom as _mom  # noqa: E402  (inline formatting helpers)

# ------------------------------------------------------------------------------------------------ palette
NAVY, NAVY2 = "#0B1030", "#141B4D"
VIOLET, BLUE, TEAL, GREEN, AMBER, RED, PINK = "#6D4AFF", "#2563EB", "#0F9D8F", "#15803D", "#B45309", "#B91C1C", "#BE185D"
INK, MUTED, TINT, LINE = "#1B2142", "#5B648F", "#F4F3FF", "#D9DCF3"
CALLOUT = {  # accent, background, label
    "NOTE": (BLUE, "#EAF1FF", "NOTE"),
    "ASSUMPTION": (AMBER, "#FFF6E3", "ASSUMPTION"),
    "DECISION": (GREEN, "#E8F8EF", "DESIGN DECISION"),
    "KEY": (VIOLET, "#F0ECFF", "KEY IDEA"),
    "CAUTION": (RED, "#FDECEC", "CAUTION"),
    "TIP": (TEAL, "#E5F7F5", "TIP"),
}


# ------------------------------------------------------------------------------------------------ parsing
@dataclass
class Block:
    kind: str
    data: object = None
    extra: object = None


def parse(md: str) -> list[Block]:
    lines = md.replace("\r", "").split("\n")
    out: list[Block] = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
        elif ln.strip() == "<pagebreak>":
            out.append(Block("pagebreak")); i += 1
        elif ln.strip().startswith("```"):
            lang = ln.strip()[3:].strip()
            i += 1
            body = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i]); i += 1
            i += 1
            out.append(Block("stats" if lang == "stats" else "code", body))
        elif ln.startswith("# "):
            out.append(Block("h1", ln[2:].strip())); i += 1
        elif ln.startswith("## "):
            title = ln[3:].strip()
            i += 1
            lead = ""
            if i < len(lines) and lines[i].startswith("LEAD:"):
                lead = lines[i][5:].strip(); i += 1
            out.append(Block("h2", title, lead))
        elif ln.startswith("### "):
            out.append(Block("h3", ln[4:].strip())); i += 1
        elif ln.startswith("#### "):
            out.append(Block("h4", ln[5:].strip())); i += 1
        elif re.match(r"^---+\s*$", ln):
            out.append(Block("hr")); i += 1
        elif re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", ln):
            m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", ln)
            out.append(Block("img", m.group(1), m.group(2))); i += 1
        elif ln.startswith(">"):
            body = []
            while i < len(lines) and lines[i].startswith(">"):
                body.append(re.sub(r"^>\s?", "", lines[i])); i += 1
            m = re.match(r"^\[!(\w+)\]\s*(.*)$", body[0])
            kind, title = (m.group(1).upper(), m.group(2).strip()) if m else ("NOTE", "")
            text = body[1:] if m else body
            out.append(Block("callout", (kind if kind in CALLOUT else "NOTE", title), "\n".join(text).strip()))
        elif ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            out.append(Block("table", rows))
        elif re.match(r"^\s*[-*]\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip()); i += 1
                while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*[-*]\s+", lines[i]):
                    items[-1] += " " + lines[i].strip(); i += 1
            out.append(Block("ul", items))
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip()); i += 1
                while i < len(lines) and lines[i].startswith("   ") and lines[i].strip() and not re.match(r"^\s*\d+[.)]\s+", lines[i]):
                    items[-1] += " " + lines[i].strip(); i += 1
            out.append(Block("ol", items))
        else:
            para = ln
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||>|```|!\[|<pagebreak>|\s*[-*]\s|\s*\d+[.)]\s|---)", lines[i]):
                sep = "\n" if para.endswith("  ") else " "
                para = para.rstrip() + sep + lines[i].lstrip()
                i += 1
            out.append(Block("p", para.strip()))
    return out


def split_number(title: str) -> tuple[str | None, str]:
    title = re.sub(r"[`*]", "", title)
    m = re.match(r"^(\d+)\.\s+(.*)$", title)
    return (m.group(1), m.group(2)) if m else (None, title)


# ================================================================================================ PDF
def _register_fonts():
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    win = Path("C:/Windows/Fonts")
    import matplotlib

    dj = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    sets = [("SegoeUI", win / "segoeui.ttf", win / "segoeuib.ttf", win / "segoeuii.ttf", win / "segoeuiz.ttf"),
            ("DejaVu", dj / "DejaVuSans.ttf", dj / "DejaVuSans-Bold.ttf", dj / "DejaVuSans-Oblique.ttf", dj / "DejaVuSans-BoldOblique.ttf")]
    for fam, r, b, i, bi in sets:
        if all(p.exists() for p in (r, b, i, bi)):
            pdfmetrics.registerFont(TTFont(fam, str(r)))
            pdfmetrics.registerFont(TTFont(fam + "-B", str(b)))
            pdfmetrics.registerFont(TTFont(fam + "-I", str(i)))
            pdfmetrics.registerFont(TTFont(fam + "-BI", str(bi)))
            pdfmetrics.registerFontFamily(fam, normal=fam, bold=fam + "-B", italic=fam + "-I", boldItalic=fam + "-BI")
            mono = "Courier"
            for cand in (win / "consola.ttf", dj / "DejaVuSansMono.ttf"):
                if cand.exists():
                    pdfmetrics.registerFont(TTFont("MonoX", str(cand)))
                    mono = "MonoX"
                    break
            return fam, mono
    return "Helvetica", "Courier"


def build_pdf(md: str, cover: dict, base_dir: Path, out_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Flowable, Frame, Image as RLImage, KeepTogether, ListFlowable, ListItem,
                                    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Preformatted, Spacer, Table, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents

    F, MONO = _register_fonts()
    FB = F + "-B" if F != "Helvetica" else "Helvetica-Bold"
    FI = F + "-I" if F != "Helvetica" else "Helvetica-Oblique"
    W, H = A4
    LM = RM = 18 * mm
    TM, BM = 22 * mm, 20 * mm
    AW = W - LM - RM
    safe = _mom._pdf_safe if F == "Helvetica" else (lambda s: s)

    body = ParagraphStyle("body", fontName=F, fontSize=9.8, leading=14.6, textColor=HexColor(INK), spaceAfter=5)
    small = ParagraphStyle("small", parent=body, fontSize=8.6, leading=12)
    cellst = ParagraphStyle("cell", parent=body, fontSize=8.7, leading=11.8, spaceAfter=0)
    cellh = ParagraphStyle("cellh", parent=cellst, fontName=FB, textColor=colors.white)
    caption = ParagraphStyle("caption", parent=body, fontName=FI, fontSize=8.4, textColor=HexColor(MUTED), alignment=1, spaceBefore=3, spaceAfter=10)
    h4 = ParagraphStyle("h4", parent=body, fontName=FB, fontSize=10.8, textColor=HexColor(VIOLET), spaceBefore=8, spaceAfter=3)
    CODE_FMT = f'<font name="{MONO}" size="8.6" color="{VIOLET}">{{}}</font>'

    def inl(text: str, code_fmt: str = CODE_FMT) -> str:
        s = _mom._inline(text, code_fmt, r"<b>\1</b>", r"<i>\1</i>", f'<link href="{{1}}" color="{BLUE}">{{0}}</link>')
        return safe(s).replace("\n", "<br/>")

    def para(text, style=body):
        return Paragraph(inl(text), style)

    # ------------------------------------------------------------ custom flowables
    class SectionBanner(Flowable):
        def __init__(self, number, title, lead, width):
            super().__init__()
            self.number, self.title, self.lead_text, self.width_ = number, title, lead, width
            self.lead_style = ParagraphStyle("lead", fontName=F, fontSize=9.6, leading=13.4, textColor=HexColor("#C9D0FF"))
            self.lead = Paragraph(inl(lead, f'<font name="{MONO}" size="8.6" color="#E2D9FF">{{}}</font>'), self.lead_style) if lead else None

        def wrap(self, aw, ah):
            self.width_ = aw
            self.lead_h = 0
            if self.lead:
                _, self.lead_h = self.lead.wrap(aw - 34 * mm, 500)
            self.height_ = 20 * mm + (self.lead_h + 3 * mm if self.lead else 0)
            return aw, self.height_

        def draw(self):
            c = self.canv
            w, h = self.width_, self.height_
            c.saveState()
            p = c.beginPath()
            p.roundRect(0, 0, w, h, 4 * mm)
            c.clipPath(p, stroke=0, fill=0)
            c.setFillColor(HexColor(NAVY)); c.rect(0, 0, w, h, fill=1, stroke=0)
            for cx, cy, r, col, a in ((w - 8 * mm, h + 6 * mm, 38 * mm, VIOLET, 0.55), (w - 45 * mm, -12 * mm, 30 * mm, BLUE, 0.35), (6 * mm, h, 26 * mm, VIOLET, 0.25)):
                c.setFillColor(HexColor(col)); c.setFillAlpha(a * 0.5); c.circle(cx, cy, r, fill=1, stroke=0)
            c.setFillAlpha(1)
            c.setFillColor(HexColor(VIOLET)); c.setStrokeColor(HexColor("#A78BFA")); c.setLineWidth(1.2)
            if self.number:
                cy = h - 11.5 * mm
                c.circle(13 * mm, cy, 6.2 * mm, fill=1, stroke=1)
                c.setFillColor(colors.white); c.setFont(FB, 14)
                c.drawCentredString(13 * mm, cy - 2 * mm, self.number)
                tx = 24 * mm
            else:
                tx = 8 * mm
            c.setFillColor(colors.white); c.setFont(FB, 19)
            c.drawString(tx, h - 13.4 * mm, safe(self.title))
            if self.lead:
                self.lead.drawOn(c, tx, h - 17 * mm - self.lead_h)
            c.restoreState()

    class SubHeading(Table):
        def __init__(self, text):
            st = ParagraphStyle("h3", fontName=FB, fontSize=12.6, leading=16, textColor=HexColor(NAVY))
            super().__init__([["", Paragraph(inl(text), st)]], colWidths=[1.6 * mm, AW - 1.6 * mm])
            self.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), HexColor(VIOLET)), ("LEFTPADDING", (1, 0), (1, 0), 7), ("TOPPADDING", (0, 0), (-1, -1), 2),
                                      ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            self.plain = re.sub(r"[*`_]", "", text)

    class Doc(BaseDocTemplate):
        def afterFlowable(self, fl):
            if isinstance(fl, SectionBanner):
                key = f"sec{self.page}_{id(fl)}"
                label = (fl.number + ".  " if fl.number else "") + fl.title
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(label, key, 0, 0)
                if fl.number:
                    self.notify("TOCEntry", (0, label, self.page - 1))
            elif isinstance(fl, SubHeading):
                key = f"sub{self.page}_{id(fl)}"
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(fl.plain, key, 1, 0)

    # ------------------------------------------------------------ page decorations
    def draw_cover(c, d):
        c.saveState()
        steps = 60
        for k in range(steps):
            t = k / steps
            r = int(7 + (24 - 7) * t); g = int(10 + (20 - 10) * t); b = int(28 + (72 - 28) * t)
            c.setFillColorRGB(r / 255, g / 255, b / 255)
            c.rect(0, H * (1 - (k + 1) / steps), W, H / steps + 1, fill=1, stroke=0)
        for cx, cy, r, col, a in ((W * 0.15, H * 0.9, 120 * mm, "#7C3AED", 0.22), (W * 0.95, H * 0.12, 110 * mm, "#2563EB", 0.22), (W * 0.7, H * 0.55, 80 * mm, "#6D4AFF", 0.10)):
            c.setFillColor(HexColor(col)); c.setFillAlpha(a); c.circle(cx, cy, r, fill=1, stroke=0)
        c.setFillAlpha(1)
        # logo mark
        x0, y0, s = 22 * mm, H - 48 * mm, 20 * mm
        c.setFillColor(HexColor(VIOLET)); c.roundRect(x0, y0, s, s, 5 * mm, fill=1, stroke=0)
        c.setFillColor(HexColor(BLUE)); c.setFillAlpha(0.55); c.roundRect(x0 + s * 0.45, y0, s * 0.55, s, 5 * mm, fill=1, stroke=0); c.setFillAlpha(1)
        c.setStrokeColor(colors.white); c.setLineWidth(1.8); c.setLineCap(1)
        for k, hh in enumerate((3, 8, 13, 8, 3)):
            xx = x0 + s * (0.2 + k * 0.15)
            c.line(xx, y0 + s / 2 - hh * 0.5 * mm * 0.62, xx, y0 + s / 2 + hh * 0.5 * mm * 0.62)
        c.setFillColor(colors.white); c.setFont(FB, 17); c.drawString(x0 + s + 5 * mm, y0 + 8.4 * mm, "MeetingAI")
        c.setFillColor(HexColor("#A9B4F5")); c.setFont(F, 8); c.drawString(x0 + s + 5 * mm, y0 + 3.6 * mm, "MEETING INTELLIGENCE")
        # title block
        c.setFillColor(HexColor("#A78BFA")); c.setFont(FB, 11); c.drawString(22 * mm, H - 92 * mm, d.get("kicker", "").upper())
        c.setFillColor(colors.white); c.setFont(FB, 44)
        y = H - 110 * mm
        for line in d["title_lines"]:
            c.drawString(22 * mm, y, line); y -= 17 * mm
        c.setFillColor(HexColor("#C9D0FF")); c.setFont(F, 12.5)
        for line in d.get("tagline_lines", []):
            c.drawString(22 * mm, y - 2 * mm, safe(line)); y -= 7 * mm
        # accent line
        c.setStrokeColor(HexColor(VIOLET)); c.setLineWidth(3); c.line(22 * mm, y - 4 * mm, 62 * mm, y - 4 * mm)
        # screenshot card
        img = d.get("image")
        if img and Path(img).exists():
            iw, ih = PILImage.open(img).size
            cw = W - 44 * mm
            chh = cw * ih / iw
            chh = min(chh, 64 * mm)
            cy = 58 * mm
            c.setFillColor(colors.black); c.setFillAlpha(0.35); c.roundRect(22 * mm + 1.5 * mm, cy - 1.8 * mm, cw, chh, 3 * mm, fill=1, stroke=0); c.setFillAlpha(1)
            p = c.beginPath(); p.roundRect(22 * mm, cy, cw, chh, 3 * mm)
            c.saveState(); c.clipPath(p, stroke=0, fill=0)
            c.drawImage(img, 22 * mm, cy - (cw * ih / iw - chh), width=cw, height=cw * ih / iw, mask="auto")
            c.restoreState()
            c.setStrokeColor(HexColor("#7C6CFF")); c.setLineWidth(1.2); c.roundRect(22 * mm, cy, cw, chh, 3 * mm, fill=0, stroke=1)
        # meta strip
        yy = 46 * mm
        c.setFillColor(HexColor("#A9B4F5")); c.setFont(F, 8)
        xcol = 22 * mm
        colw = (W - 44 * mm) / max(1, len(d["meta"]))
        for label, value in d["meta"]:
            c.setFillColor(HexColor("#8E99E0")); c.setFont(FB, 7.6); c.drawString(xcol, yy, label.upper())
            c.setFillColor(colors.white); c.setFont(F, 8.8)
            txt = value.replace("https://", "").rstrip("/") if value.startswith("http") else value
            fs = 8.8
            while c.stringWidth(txt, F, fs) > colw - 4 * mm and fs > 5.8:
                fs -= 0.2
            c.setFont(F, fs)
            c.drawString(xcol, yy - 5.2 * mm, txt)
            if value.startswith("http"):
                c.linkURL(value, (xcol, yy - 7 * mm, xcol + colw - 4 * mm, yy + 3 * mm), relative=0)
            xcol += colw
        c.setFillColor(HexColor("#6F7BC4")); c.setFont(F, 7.6)
        c.drawString(22 * mm, 16 * mm, d.get("footer", ""))
        c.restoreState()

    def draw_body(c, d):
        c.saveState()
        n = d.page - 1
        # top accent bar (violet -> blue)
        steps = 40
        for k in range(steps):
            t = k / steps
            r = (0x6D + (0x25 - 0x6D) * t) / 255; g = (0x4A + (0x63 - 0x4A) * t) / 255; b = (0xFF + (0xEB - 0xFF) * t) / 255
            c.setFillColorRGB(r, g, b)
            c.rect(W * k / steps, H - 3.2 * mm, W / steps + 0.5, 3.2 * mm, fill=1, stroke=0)
        c.setFont(FB, 8); c.setFillColor(HexColor(VIOLET)); c.drawString(LM, H - 11 * mm, "MeetingAI")
        c.setFont(F, 8); c.setFillColor(HexColor(MUTED)); c.drawString(LM + 17 * mm, H - 11 * mm, "·  " + cover["doc_title"])
        c.setStrokeColor(HexColor(LINE)); c.setLineWidth(0.6); c.line(LM, H - 14 * mm, W - RM, H - 14 * mm)
        c.line(LM, 14 * mm, W - RM, 14 * mm)
        c.setFont(F, 7.8); c.setFillColor(HexColor(MUTED)); c.drawString(LM, 9.3 * mm, cover.get("footer_left", ""))
        c.setFillColor(HexColor(NAVY)); c.roundRect(W - RM - 14 * mm, 7 * mm, 14 * mm, 5.6 * mm, 2.8 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white); c.setFont(FB, 8.2); c.drawCentredString(W - RM - 7 * mm, 8.8 * mm, str(n))
        c.restoreState()

    doc = Doc(str(out_path), pagesize=A4, leftMargin=LM, rightMargin=RM, topMargin=TM, bottomMargin=BM,
              title=f"MeetingAI - {cover['doc_title']}", author="MeetingAI", subject=cover.get("tagline", ""))
    frame_body = Frame(LM, BM, AW, H - TM - BM, id="body", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    frame_cover = Frame(0, 0, W, H, id="cover")
    doc.addPageTemplates([PageTemplate(id="cover", frames=[frame_cover], onPage=lambda c, d: draw_cover(c, cover)),
                          PageTemplate(id="body", frames=[frame_body], onPage=lambda c, d: draw_body(c, d))])

    # ------------------------------------------------------------ flowable builders
    fig_no = [0]

    def make_table(rows):
        head, *data = rows
        ncol = len(head)
        weights = []
        for j in range(ncol):
            longest = max([len(re.sub(r"[*_`]", "", head[j]))] + [len(re.sub(r"[*_`\[\]]", "", r[j])) for r in data if j < len(r)])
            weights.append(min(max(longest, 4), 70) ** 0.78)
        tot = sum(weights)
        widths = [max(AW * x / tot, 14 * mm) for x in weights]
        s = sum(widths)
        widths = [x * AW / s for x in widths]
        cells = [[Paragraph(inl(h), cellh) for h in head]]
        for r in data:
            cells.append([Paragraph(inl(r[j] if j < len(r) else ""), cellst) for j in range(ncol)])
        t = Table(cells, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), HexColor(NAVY)), ("LINEBELOW", (0, 0), (-1, 0), 1.6, HexColor(VIOLET)),
                 ("GRID", (0, 0), (-1, -1), 0.4, HexColor(LINE)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("LEFTPADDING", (0, 0), (-1, -1), 5.5), ("RIGHTPADDING", (0, 0), (-1, -1), 5.5), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, HexColor("#F7F6FF")])]
        t.setStyle(TableStyle(style))
        return t

    def make_callout(kind, title, text):
        accent, bg, label = CALLOUT[kind]
        lab = f'<font name="{FB}" size="7.8" color="{accent}">{label}{(" · " + _mom._html.escape(title).upper()) if title else ""}</font>'
        inner = [Paragraph(lab, ParagraphStyle("cl", fontName=FB, fontSize=7.8, leading=10, spaceAfter=3))]
        for chunk in re.split(r"\n\s*\n", text):
            lines = [l for l in chunk.split("\n") if l.strip()]
            if lines and all(re.match(r"^\s*[-*]\s+", l) for l in lines):
                inner.append(ListFlowable([ListItem(para(re.sub(r"^\s*[-*]\s+", "", l), cellst), leftIndent=11) for l in lines], bulletType="bullet", start="•", bulletColor=HexColor(accent),
                                          bulletFontSize=9, leftIndent=11))
            elif chunk.strip():
                inner.append(para(chunk.strip().replace("\n", " "), ParagraphStyle("cb", parent=body, fontSize=9.2, leading=13.4, spaceAfter=3)))
        t = Table([[inner]], colWidths=[AW])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor(bg)), ("LINEBEFORE", (0, 0), (0, -1), 3.2, HexColor(accent)),
                               ("BOX", (0, 0), (-1, -1), 0.4, HexColor(accent)), ("LEFTPADDING", (0, 0), (-1, -1), 11), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                               ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        return t

    def make_stats(rows):
        items = [[x.strip() for x in r.split("|", 1)] for r in rows if "|" in r]
        n = len(items)
        gap = 0
        cw = AW / n
        vst = ParagraphStyle("sv", fontName=FB, fontSize=19, leading=22, textColor=HexColor(VIOLET), alignment=1)
        lst = ParagraphStyle("sl", fontName=F, fontSize=8.2, leading=10.5, textColor=HexColor(MUTED), alignment=1)
        cells, widths = [], []
        for k, (v, l) in enumerate(items):
            cells.append([Paragraph(inl(v), vst), Paragraph(inl(l), lst)])
        row = [[c] for c in cells]
        data = [[Table([[c[0]], [c[1]]], colWidths=[cw]) for c in cells]]
        t = Table(data, colWidths=[cw] * n, hAlign="LEFT")
        sty = [("BACKGROUND", (i, 0), (i, 0), HexColor(TINT)) for i in range(n)] + [("BOX", (i, 0), (i, 0), 0.6, HexColor(LINE)) for i in range(n)]
        sty += [("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]
        t.setStyle(TableStyle(sty))
        return t

    def make_code(lines):
        txt = safe("\n".join(lines))
        pre = Preformatted(txt, ParagraphStyle("pre", fontName=MONO, fontSize=8, leading=10.4, textColor=HexColor("#DCE2FF")))
        t = Table([[pre]], colWidths=[AW])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor(NAVY)), ("LINEBEFORE", (0, 0), (0, -1), 3, HexColor(VIOLET)), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                               ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
        return t

    def make_figure(alt, rel):
        path = base_dir / rel
        if not path.exists():
            return [para(f"[missing image: {rel}]")]
        fig_no[0] += 1
        iw, ih = PILImage.open(path).size
        maxw, maxh = AW, 215 * mm
        sc = min(maxw / iw, maxh / ih)
        img = RLImage(str(path), width=iw * sc, height=ih * sc)
        is_shot = path.suffix.lower() in (".jpg", ".jpeg")
        if is_shot:
            tb = Table([[img]], colWidths=[iw * sc + 4])
            tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, HexColor("#B9BFE8")), ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                                    ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            img = tb
        cap = Paragraph(f'<font name="{FB}" color="{VIOLET}">Figure {fig_no[0]}</font> &nbsp;·&nbsp; {inl(alt)}', caption)
        return [KeepTogether([img, cap])]

    # ------------------------------------------------------------ assemble story
    blocks = parse(md)
    story: list = [Spacer(1, 1), NextPageTemplate("body"), PageBreak()]
    # table of contents
    toc = TableOfContents()
    toc.dotsMinLevel = 0
    toc.levelStyles = [ParagraphStyle("toc0", fontName=FB, fontSize=10.4, leading=19, textColor=HexColor(NAVY), leftIndent=4, rightIndent=14, firstLineIndent=0)]
    story.append(Paragraph('<font color="%s">Contents</font>' % NAVY, ParagraphStyle("tch", fontName=FB, fontSize=24, leading=30, spaceAfter=4)))
    story.append(Table([[""]], colWidths=[24 * mm], rowHeights=[1.6 * mm], style=[("BACKGROUND", (0, 0), (-1, -1), HexColor(VIOLET))], hAlign="LEFT"))
    story += [Spacer(1, 6 * mm), toc, PageBreak()]
    first_section = True
    for b in blocks:
        k = b.kind
        if k == "h1":
            continue
        if k == "pagebreak":
            story.append(PageBreak())
        elif k == "h2":
            num, title = split_number(b.data)
            story += [CondPageBreak(95 * mm) if not first_section else Spacer(1, 0), Spacer(1, 2 * mm), SectionBanner(num, title, b.extra, AW), Spacer(1, 4.5 * mm)]
            first_section = False
        elif k == "h3":
            story += [Spacer(1, 3 * mm), CondPageBreak(40 * mm), SubHeading(b.data), Spacer(1, 2.5 * mm)]
        elif k == "h4":
            story.append(para(b.data, h4))
        elif k == "p":
            story.append(para(b.data))
        elif k == "ul":
            story.append(ListFlowable([ListItem(para(x), leftIndent=13) for x in b.data], bulletType="bullet", start="•", bulletColor=HexColor(VIOLET), bulletFontSize=10, leftIndent=13))
            story.append(Spacer(1, 3))
        elif k == "ol":
            story.append(ListFlowable([ListItem(para(x), leftIndent=15) for x in b.data], bulletType="1", bulletFontName=FB, bulletColor=HexColor(VIOLET), bulletFontSize=9.4, leftIndent=15))
            story.append(Spacer(1, 3))
        elif k == "table":
            story += [make_table(b.data), Spacer(1, 5 * mm)]
        elif k == "callout":
            story += [Spacer(1, 1.5 * mm), make_callout(b.data[0], b.data[1], b.extra), Spacer(1, 4 * mm)]
        elif k == "stats":
            story += [make_stats(b.data), Spacer(1, 5 * mm)]
        elif k == "code":
            story += [make_code(b.data), Spacer(1, 4 * mm)]
        elif k == "img":
            story += make_figure(b.data, b.extra)
        elif k == "hr":
            story.append(Spacer(1, 2 * mm))
    doc.multiBuild(story)


# ================================================================================================ DOCX
def build_docx(md: str, cover: dict, base_dir: Path, out_path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    FONT = "Segoe UI"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.9)
    text_w = Cm(21 - 3.8)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10)
    st.font.color.rgb = rgb(INK)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.18
    for name, size, col in (("Heading 1", 20, "#FFFFFF"), ("Heading 2", 13.5, NAVY), ("Heading 3", 11.5, VIOLET)):
        hs = doc.styles[name]
        hs.font.name = FONT
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = rgb(col)
        hs.paragraph_format.space_before = Pt(10)
        hs.paragraph_format.space_after = Pt(4)

    # ------------------------------------------------------------ helpers
    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcol.lstrip("#"))
        tcPr.append(shd)

    def borders(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            spec = kw.get(edge)
            el = OxmlElement(f"w:{edge}")
            if spec:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:color"), spec[1].lstrip("#")); el.set(qn("w:space"), "0")
            else:
                el.set(qn("w:val"), "nil")
            tb.append(el)
        tcPr.append(tb)

    def cell_margins(cell, top=60, bottom=60, left=110, right=110):
        tcPr = cell._tc.get_or_add_tcPr()
        m = OxmlElement("w:tcMar")
        for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
        tcPr.append(m)

    def set_width(table, widths):
        table.autofit = False
        for row in table.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = w

    def no_borders(table):
        tbl = table._tbl
        tblPr = tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); b.append(e)
        tblPr.append(b)

    def runs(par, text, size=None, color=None, bold=None):
        before = len(par.runs)
        _mom._add_runs(par, text)
        for r in par.runs[before:]:
            r.font.name = r.font.name or FONT
            if size:
                r.font.size = Pt(size)
            if color and not r.font.color.rgb:
                r.font.color.rgb = rgb(color)
            if bold is not None:
                r.bold = bold or r.bold

    def field(par, instr):
        r = par.add_run()
        for t, txt in (("begin", None), (None, instr), ("end", None)):
            if t:
                f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        return r

    # ------------------------------------------------------------ cover
    sec.different_first_page_header_footer = True
    cv = doc.add_table(rows=1, cols=1)
    cv.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = cv.rows[0].cells[0]
    shade(c0, "#0B1030")
    cell_margins(c0, 500, 400, 520, 520)
    no_borders(cv)
    cv.rows[0].height = Cm(19.5)
    p = c0.paragraphs[0]
    r = p.add_run("MeetingAI"); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = rgb("#FFFFFF"); r.font.name = FONT
    p = c0.add_paragraph(); r = p.add_run("MEETING INTELLIGENCE"); r.font.size = Pt(8); r.font.color.rgb = rgb("#A9B4F5"); r.font.name = FONT
    for _ in range(3):
        c0.add_paragraph()
    p = c0.add_paragraph(); r = p.add_run(cover.get("kicker", "").upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = rgb("#A78BFA"); r.font.name = FONT
    for line in cover["title_lines"]:
        p = c0.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line); r.bold = True; r.font.size = Pt(38); r.font.color.rgb = rgb("#FFFFFF"); r.font.name = FONT
    p = c0.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    r = p.add_run(" ".join(cover.get("tagline_lines", []))); r.font.size = Pt(12); r.font.color.rgb = rgb("#C9D0FF"); r.font.name = FONT
    img = cover.get("image")
    if img and Path(img).exists():
        p = c0.add_paragraph(); p.paragraph_format.space_before = Pt(18)
        p.add_run().add_picture(str(img), width=Cm(15.8))
    for label, value in cover["meta"]:
        p = c0.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(5)
        r = p.add_run(label.upper() + "   "); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = rgb("#8E99E0"); r.font.name = FONT
        if value.startswith("http"):
            _mom._add_hyperlink(p, value, value)
            for rr in p.runs[-1:]:
                pass
        else:
            r = p.add_run(value); r.font.size = Pt(9); r.font.color.rgb = rgb("#FFFFFF"); r.font.name = FONT
    doc.add_page_break()

    # ------------------------------------------------------------ header / footer
    hdr = sec.header.paragraphs[0]
    r = hdr.add_run("MeetingAI"); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = rgb(VIOLET); r.font.name = FONT
    r = hdr.add_run("  ·  " + cover["doc_title"]); r.font.size = Pt(8); r.font.color.rgb = rgb(MUTED); r.font.name = FONT
    pPr = hdr._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6"); bt.set(qn("w:color"), "D9DCF3"); bt.set(qn("w:space"), "4")
    bd.append(bt); pPr.append(bd)
    ft = sec.footer.paragraphs[0]
    ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = ft.add_run(cover.get("footer_left", "") + "      Page "); r.font.size = Pt(8); r.font.color.rgb = rgb(MUTED); r.font.name = FONT
    fr = field(ft, "PAGE"); fr.font.size = Pt(8); fr.bold = True; fr.font.color.rgb = rgb(VIOLET)

    # ------------------------------------------------------------ content
    blocks = parse(md)
    sections = [split_number(b.data) for b in blocks if b.kind == "h2" and split_number(b.data)[0]]
    # contents
    cp = doc.add_paragraph(); r = cp.add_run("Contents"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = rgb(NAVY); r.font.name = FONT
    for num, title in sections:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(num.rjust(2) + "   "); r.bold = True; r.font.color.rgb = rgb(VIOLET); r.font.size = Pt(10.5); r.font.name = FONT
        r = p.add_run(title); r.font.size = Pt(10.5); r.font.color.rgb = rgb(NAVY); r.font.name = FONT
    p = doc.add_paragraph(); r = p.add_run("Tip: open the Navigation Pane (View → Navigation) to jump between sections."); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = rgb(MUTED)
    doc.add_page_break()

    fig_no = [0]
    first = True

    def callout_table(kind, title, text):
        accent, bg, label = CALLOUT[kind]
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.rows[0].cells[0]
        shade(c, bg); borders(c, left=(24, accent)); cell_margins(c, 90, 80, 200, 160)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + (" · " + title.upper() if title else "")); r.bold = True; r.font.size = Pt(7.8); r.font.color.rgb = rgb(accent); r.font.name = FONT
        for chunk in re.split(r"\n\s*\n", text):
            lines = [l for l in chunk.split("\n") if l.strip()]
            if lines and all(re.match(r"^\s*[-*]\s+", l) for l in lines):
                for l in lines:
                    q = c.add_paragraph(); q.paragraph_format.left_indent = Cm(0.45); q.paragraph_format.first_line_indent = Cm(-0.35); q.paragraph_format.space_after = Pt(1.5)
                    rr = q.add_run("•  "); rr.font.color.rgb = rgb(accent)
                    runs(q, re.sub(r"^\s*[-*]\s+", "", l), 9.4)
            elif chunk.strip():
                q = c.add_paragraph(); q.paragraph_format.space_after = Pt(2)
                runs(q, chunk.strip().replace("\n", " "), 9.4)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    for b in blocks:
        k = b.kind
        if k == "h1":
            continue
        if k == "pagebreak":
            doc.add_page_break()
        elif k == "h2":
            num, title = split_number(b.data)
            if not first:
                doc.add_page_break()
            first = False
            t = doc.add_table(rows=1, cols=1)
            no_borders(t)
            c = t.rows[0].cells[0]
            shade(c, NAVY); cell_margins(c, 160, 160, 260, 260)
            p = c.paragraphs[0]
            p.style = doc.styles["Heading 1"]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
            if num:
                rn = p.add_run(num + "   "); rn.font.color.rgb = rgb("#A78BFA"); rn.font.size = Pt(20); rn.bold = True
            rt = p.add_run(title); rt.font.color.rgb = rgb("#FFFFFF"); rt.font.size = Pt(20); rt.bold = True
            if b.extra:
                q = c.add_paragraph(); q.paragraph_format.space_after = Pt(0)
                runs(q, b.extra, 10, "#C9D0FF")
                for rr in q.runs:
                    rr.font.color.rgb = rgb("#C9D0FF")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == "h3":
            p = doc.add_paragraph(style="Heading 2")
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); bl = OxmlElement("w:left")
            bl.set(qn("w:val"), "single"); bl.set(qn("w:sz"), "24"); bl.set(qn("w:color"), VIOLET.lstrip("#")); bl.set(qn("w:space"), "6")
            bd.append(bl); pPr.append(bd)
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.keep_with_next = True
            p.add_run(re.sub(r"[*`_]", "", b.data))
        elif k == "h4":
            p = doc.add_paragraph(style="Heading 3"); p.add_run(re.sub(r"[*`_]", "", b.data))
        elif k == "p":
            p = doc.add_paragraph(); runs(p, b.data)
        elif k in ("ul", "ol"):
            for n, item in enumerate(b.data, 1):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.5); p.paragraph_format.space_after = Pt(2.5)
                r = p.add_run(("•" if k == "ul" else f"{n}.") + "  "); r.bold = True; r.font.color.rgb = rgb(VIOLET)
                runs(p, item)
        elif k == "table":
            head, *data = b.data
            ncol = len(head)
            weights = []
            for j in range(ncol):
                longest = max([len(re.sub(r"[*_`]", "", head[j]))] + [len(re.sub(r"[*_`\[\]]", "", r[j])) for r in data if j < len(r)])
                weights.append(min(max(longest, 4), 70) ** 0.78)
            tot = sum(weights)
            widths = [int(text_w * w / tot) for w in weights]
            t = doc.add_table(rows=1 + len(data), cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_width(t, widths)
            for j, h in enumerate(head):
                c = t.rows[0].cells[j]
                shade(c, NAVY); borders(c, bottom=(12, VIOLET)); cell_margins(c, 70, 70)
                pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
                rr = pp.add_run(re.sub(r"[*`_]", "", h)); rr.bold = True; rr.font.size = Pt(8.8); rr.font.color.rgb = rgb("#FFFFFF"); rr.font.name = FONT
            for i, row in enumerate(data, 1):
                for j in range(ncol):
                    c = t.rows[i].cells[j]
                    shade(c, "#F7F6FF" if i % 2 == 0 else "#FFFFFF")
                    borders(c, top=(4, LINE), bottom=(4, LINE), left=(4, LINE), right=(4, LINE)); cell_margins(c)
                    pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
                    runs(pp, row[j] if j < len(row) else "", 8.8)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif k == "callout":
            callout_table(b.data[0], b.data[1], b.extra)
        elif k == "stats":
            items = [[x.strip() for x in r_.split("|", 1)] for r_ in b.data if "|" in r_]
            t = doc.add_table(rows=1, cols=len(items))
            no_borders(t)
            for j, (v, l) in enumerate(items):
                c = t.rows[0].cells[j]
                shade(c, TINT); borders(c, left=(8, "#FFFFFF"), right=(8, "#FFFFFF")); cell_margins(c, 120, 120)
                p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
                rr = p.add_run(re.sub(r"[*`_]", "", v)); rr.bold = True; rr.font.size = Pt(18); rr.font.color.rgb = rgb(VIOLET); rr.font.name = FONT
                q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after = Pt(0)
                rr = q.add_run(re.sub(r"[*`_]", "", l)); rr.font.size = Pt(8); rr.font.color.rgb = rgb(MUTED); rr.font.name = FONT
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif k == "code":
            t = doc.add_table(rows=1, cols=1)
            no_borders(t)
            c = t.rows[0].cells[0]
            shade(c, NAVY); borders(c, left=(24, VIOLET)); cell_margins(c, 100, 100, 200, 160)
            for n, line in enumerate(b.data):
                pp = c.paragraphs[0] if n == 0 else c.add_paragraph()
                pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.0
                rr = pp.add_run(line if line else " "); rr.font.name = "Consolas"; rr.font.size = Pt(8); rr.font.color.rgb = rgb("#DCE2FF")
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif k == "img":
            path = base_dir / b.extra
            if path.exists():
                fig_no[0] += 1
                pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.keep_with_next = True
                pp.add_run().add_picture(str(path), width=Cm(16.4) if path.suffix.lower() == ".png" else Cm(16.0))
                cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cp.add_run(f"Figure {fig_no[0]}"); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = rgb(VIOLET)
                r = cp.add_run("  ·  " + re.sub(r"[*`_]", "", b.data)); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = rgb(MUTED)
        elif k == "hr":
            pass
    doc.save(str(out_path))


# ================================================================================================ driver
def default_cover(doc_title: str, kicker: str, title_lines: list[str], tagline_lines: list[str], site: str, repo: str, image: str | None) -> dict:
    return {"doc_title": doc_title, "kicker": kicker, "title_lines": title_lines, "tagline_lines": tagline_lines, "image": image,
            "meta": [("Live application", site), ("Source code", repo), ("Edition", date.today().strftime("%B %Y"))],
            "footer": "MeetingAI  ·  hybrid NLP + LLM meeting intelligence  ·  Streamlit · SQLite · ChromaDB · pyannote",
            "footer_left": "MeetingAI  ·  " + doc_title}
